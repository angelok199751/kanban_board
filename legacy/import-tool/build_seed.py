#!/usr/bin/env python3
"""
QA Flow: импорт из Google Sheets (xlsx) в Supabase.

Читает выгруженную из Google Sheets книгу и создаёт ОДИН SQL-файл. Его нужно вставить в
Supabase → SQL Editor и выполнить (после миграций 0001–0003). Ключи доступа не нужны.

Свойства импорта:
  • исходный файл не изменяется;
  • повторный запуск не создаёт дублей (ключи: boards.legacy_key, tasks.legacy_id —
    это старый столбец _id) и по умолчанию НЕ перезаписывает уже существующие записи,
    чтобы правки, сделанные в новом приложении, не затирались. Флаг --overwrite меняет это;
  • весь импорт — одна транзакция: при ошибке не остаётся половины данных;
  • в конце записывается отчёт в таблицу import_reports (см. docs/import.md).

Файл с данными в репозиторий НЕ коммитится (см. .gitignore).

Пример:
  python3 tools/import/build_seed.py "Таблица.xlsx" \
      --board-sheet "143 релиз" --knowledge-sheet "База знаний" \
      --out tools/import/out/seed.sql
"""
import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl

# --- те же псевдонимы заголовков, что и в старом Code.gs -----------------------------------
ALIASES = {
    "номер": ["номер", "id", "№", "№ п/п", "номер задачи"],
    "задача": ["задача", "название", "task"],
    "статус": ["статус", "status"],
    "исполнитель": ["исполнитель", "qa", "тестировщик"],
    "разраб": ["разраб", "разработчик", "developer"],
    "дата выкатки": ["дата выкатки", "выкатка", "релиз"],
    "категория": ["категория", "тип"],
    "флоу": ["флоу", "flow", "проект"],
    "версия": ["версия", "version"],
    "коммент": ["коммент", "комментарий", "comment"],
}
SELECT_KEYS = ["исполнитель", "разраб", "категория", "флоу"]
PLAIN_KEYS = ["исполнитель", "разраб", "категория", "флоу", "версия"]  # лежат в props
COLORS = ["blue", "green", "orange", "purple", "pink", "yellow", "red", "gray"]
ARCHIVE_STATUS = "архив"
MOSCOW = "+03"  # часовой пояс таблицы неизвестен; берём московский (как у пользователей)

B62 = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz"


def norm(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "").strip().lower())


def frac_key(i):
    """Ключи порядка, совместимые с npm-пакетом fractional-indexing (generateNKeysBetween)."""
    if i < 62:
        return "a" + B62[i]
    if i < 62 * 62:
        return "b" + B62[i // 62] + B62[i % 62]
    raise ValueError("слишком много карточек в одном столбце")


def clean_text(v):
    """Значение ячейки → строка. 143.0 → '143', даты → ISO."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0, 0) else v.isoformat(sep=" ")
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def as_date(v, warn, ctx):
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", s)
    if m:
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
        except ValueError:
            pass
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", s):
        return s
    warn(f"{ctx}: не удалось распознать дату «{s}» — значение пропущено")
    return ""


def list_from_validation(ws, col_idx):
    """Значения списка из правила проверки данных столбца (только списки, заданные прямо в правиле)."""
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list" or not dv.formula1:
            continue
        f = dv.formula1.strip()
        if not (f.startswith('"') and f.endswith('"')):
            continue
        for rng in str(dv.sqref).split():
            m = re.match(r"([A-Z]+)(\d+)(?::([A-Z]+)(\d+))?$", rng)
            if not m:
                continue
            c1 = openpyxl.utils.column_index_from_string(m[1])
            c2 = openpyxl.utils.column_index_from_string(m[3] or m[1])
            if c1 <= col_idx <= c2:
                items = [x.strip() for x in f[1:-1].split(",")]
                seen, out = set(), []
                for x in items:
                    if x and x not in seen:
                        seen.add(x)
                        out.append(x)
                return out
    return None


def build(args):
    warnings = []
    warn = warnings.append

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    for name in (args.board_sheet, args.knowledge_sheet):
        if name not in wb.sheetnames:
            sys.exit(f"Лист «{name}» не найден. Есть: {', '.join(wb.sheetnames)}")

    # ---------------------------------------------------------------- доска
    ws = wb[args.board_sheet]
    headers = [c.value for c in ws[1]]
    used, cols = set(), []  # cols: dict(key, title, idx, builtin)
    id_idx = None
    for i, h in enumerate(headers):
        title = str(h).strip() if h is not None else ""
        if not title:
            continue
        if norm(title) == "_id":
            id_idx = i
            continue
        if title.startswith("_"):
            continue
        key = next((k for k, al in ALIASES.items() if k not in used and norm(title) in map(norm, al)), None)
        builtin = key is not None
        if not key:
            key = norm(title)
            n = 2
            while key in used or key in ALIASES:
                key = f"{norm(title)} {n}"
                n += 1
        used.add(key)
        cols.append({"key": key, "title": title, "idx": i, "builtin": builtin})
    by_key = {c["key"]: c for c in cols}
    for need in ("задача", "статус"):
        if need not in by_key:
            sys.exit(f"В листе «{args.board_sheet}» не найден столбец «{need}»")

    rows = []
    for r_i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        title = clean_text(row[by_key["задача"]["idx"]]).strip()
        if not title:
            if any(v not in (None, "") for v in row):
                warn(f"Строка {r_i}: есть данные, но нет названия задачи — строка пропущена")
            continue
        rows.append((r_i, row))

    # --- статусы: порядок из правила листа; «Архив» столбцом не делаем
    status_idx = by_key["статус"]["idx"] + 1
    dv_status = list_from_validation(ws, status_idx) or []
    used_status = []
    for _, row in rows:
        s = clean_text(row[by_key["статус"]["idx"]]).strip()
        if s and s not in used_status:
            used_status.append(s)
    status_order = [s for s in dv_status if norm(s) != ARCHIVE_STATUS]
    for s in used_status:
        if norm(s) != ARCHIVE_STATUS and s not in status_order:
            status_order.append(s)
            warn(f"Статус «{s}» используется в задачах, но его нет в списке листа — добавлен как столбец")

    # --- свойства
    props = []
    for pos, c in enumerate(cols, start=1):
        k = c["key"]
        if c["builtin"]:
            t = {"номер": "number", "задача": "title", "статус": "status",
                 "дата выкатки": "date", "версия": "text", "коммент": "longtext"}.get(k, "select")
        else:
            vals = [row[c["idx"]] for _, row in rows if row[c["idx"]] not in (None, "")]
            dv = list_from_validation(ws, c["idx"] + 1)
            if dv:
                t = "select"
            elif vals and all(isinstance(v, (dt.datetime, dt.date)) for v in vals):
                t = "date"
            elif vals and all(re.fullmatch(r"https?://\S+", str(v).strip()) for v in vals):
                t = "url"
            else:
                t = "text"
        props.append({"key": k, "title": c["title"], "type": t, "builtin": c["builtin"],
                      "legacy_header": c["title"], "position": pos})
    ptype = {p["key"]: p["type"] for p in props}

    # --- значения списков
    options = {}
    for p in props:
        if p["type"] != "select":
            continue
        c = by_key[p["key"]]
        vals = list_from_validation(ws, c["idx"] + 1) or []
        for _, row in rows:
            v = clean_text(row[c["idx"]]).strip()
            if v and v not in vals:
                vals.append(v)
                warn(f"«{p['title']}»: значение «{v}» используется, но не входит в список листа — добавлено")
        options[p["key"]] = [{"value": v, "color": COLORS[i % len(COLORS)], "position": i + 1}
                             for i, v in enumerate(sorted(vals, key=str.casefold))]

    # --- задачи
    nums = []
    for _, row in rows:
        n = row[by_key["номер"]["idx"]] if "номер" in by_key else None
        if n not in (None, ""):
            try:
                nums.append(int(float(n)))
            except ValueError:
                warn(f"Номер «{n}» не число — будет присвоен новый")
    seen_nums = set()
    next_num = (max(nums) if nums else 0) + 1

    per_col = {}
    tasks, assigned = [], []
    for r_i, row in rows:
        g = lambda key: row[by_key[key]["idx"]] if key in by_key else None
        status = clean_text(g("статус")).strip()
        archived = norm(status) == ARCHIVE_STATUS
        if not status:
            status = status_order[0] if status_order else "Не залито"
            warn(f"Строка {r_i}: пустой статус — поставлен «{status}»")
        elif archived:
            warn(f"Строка {r_i}: статус «Архив» — задача перенесена в архив, столбец не создаётся")
            status = status_order[0] if status_order else "Не залито"

        number = None
        raw_n = g("номер")
        if raw_n not in (None, ""):
            try:
                number = int(float(raw_n))
            except ValueError:
                number = None
        if number is None or number in seen_nums:
            if number in seen_nums:
                warn(f"Строка {r_i}: номер {number} повторяется — присвоен новый")
            number = next_num
            next_num += 1
            assigned.append(number)
        seen_nums.add(number)

        pr = {}
        for k in PLAIN_KEYS:
            if k in by_key:
                v = clean_text(g(k)).strip()
                if v:
                    pr[k] = v
        if "дата выкатки" in by_key:
            d = as_date(g("дата выкатки"), warn, f"Строка {r_i}, «Дата выкатки»")
            if d:
                pr["дата выкатки"] = d
        for c in cols:
            if c["builtin"]:
                continue
            raw = row[c["idx"]]
            if raw in (None, ""):
                continue
            pr[c["key"]] = (as_date(raw, warn, f"Строка {r_i}, «{c['title']}»")
                            if ptype[c["key"]] == "date" else clean_text(raw).strip())
            if pr[c["key"]] == "":
                pr.pop(c["key"])

        legacy = clean_text(row[id_idx]).strip() if id_idx is not None else ""
        if not legacy:
            legacy = f"row:{args.board_sheet}:{r_i}"
            warn(f"Строка {r_i}: нет _id — использован временный ключ {legacy}")

        idx = per_col.get(status, 0)
        per_col[status] = idx + 1
        tasks.append({
            "legacy_id": legacy, "number": number, "title": clean_text(g("задача")).strip(),
            "status": status, "position": frac_key(idx),
            "description": clean_text(g("коммент")) if "коммент" in by_key else "",
            "props": pr, "archived": archived,
        })

    ids = [t["legacy_id"] for t in tasks]
    if len(ids) != len(set(ids)):
        sys.exit("Дубли в столбце _id — исправьте в исходной таблице (строки скопированы).")
    if assigned:
        warn(f"Задачам без номера присвоены номера {assigned[0]}…{assigned[-1]} (всего {len(assigned)}) в порядке строк листа")

    # ------------------------------------------------------ база знаний
    wk = wb[args.knowledge_sheet]
    kh = [norm(c.value) for c in wk[1]]
    need = ["_id", "заголовок", "содержание", "метки", "обновлено"]
    if kh[:5] != need:
        sys.exit(f"Лист «{args.knowledge_sheet}» не похож на базу знаний (заголовки: {kh[:5]})")
    articles = []
    for r_i, row in enumerate(wk.iter_rows(min_row=2, values_only=True), start=2):
        title = clean_text(row[1]).strip()
        if not title:
            continue
        upd = row[4]
        if isinstance(upd, dt.datetime):
            upd_s = upd.strftime("%Y-%m-%d %H:%M:%S") + MOSCOW
        elif upd:
            m = re.match(r"(\d{2})\.(\d{2})\.(\d{4}) (\d{2}):(\d{2})", str(upd))
            upd_s = f"{m[3]}-{m[2]}-{m[1]} {m[4]}:{m[5]}:00{MOSCOW}" if m else None
        else:
            upd_s = None
        legacy = clean_text(row[0]).strip() or f"row:{args.knowledge_sheet}:{r_i}"
        articles.append({"legacy_id": legacy, "title": title,
                         "content": clean_text(row[2]), "labels": clean_text(row[3]).strip(),
                         "updated_at": upd_s})

    payload = {
        "overwrite": bool(args.overwrite),
        "board": {"legacy_key": f"sheet:{args.board_sheet}", "name": args.board_name or args.board_sheet},
        "properties": props,
        "options": options,
        "columns": [{"name": s} for s in status_order],
        "tasks": tasks,
        "knowledge": {"legacy_key": f"sheet:{args.knowledge_sheet}",
                      "name": args.knowledge_name or args.knowledge_sheet, "articles": articles},
        "source": {
            "tasks": len(tasks),
            "archived_tasks": sum(1 for t in tasks if t["archived"]),
            "articles": len(articles),
            "properties": len(props),
            "options": {k: len(v) for k, v in options.items()},
            "columns": len(status_order),
        },
        "warnings": warnings,
    }
    return payload


SQL_TEMPLATE = Path(__file__).with_name("seed_template.sql").read_text(encoding="utf-8")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsx")
    ap.add_argument("--board-sheet", default="143 релиз")
    ap.add_argument("--board-name")
    ap.add_argument("--knowledge-sheet", default="База знаний")
    ap.add_argument("--knowledge-name")
    ap.add_argument("--overwrite", action="store_true",
                    help="перезаписывать уже импортированные задачи/статьи (по умолчанию — пропускать)")
    ap.add_argument("--out", default="tools/import/out/seed.sql")
    args = ap.parse_args()

    payload = build(args)
    body = json.dumps(payload, ensure_ascii=False, separators=(",", ":"))
    assert "$qaflow_json$" not in body
    sql = SQL_TEMPLATE.replace("__PAYLOAD__", body)

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(sql, encoding="utf-8")

    s = payload["source"]
    print(f"Готово: {out}")
    print(f"  задач: {s['tasks']} (в архив: {s['archived_tasks']}), статей: {s['articles']}, "
          f"свойств: {s['properties']}, столбцов: {s['columns']}, значений списков: {s['options']}")
    if payload["warnings"]:
        print("Замечания:")
        for w in payload["warnings"]:
            print("  •", w)


if __name__ == "__main__":
    main()
